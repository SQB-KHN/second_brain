"""Synthetic dataset contract checks; never touches personal DATA_DIR."""
import concurrent.futures
import hashlib
import io
import json
import os
from pathlib import Path
import sqlite3
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile
import config
import dataset_parser as parser
from dataset_registry import Registry
from dataset_profile import profile, describe

CSV = b'number,group,date,mixed,identifier,empty\n1,A,2026-01-01,1,001,\n3,B,2026-01-02,oops,002,\n3,B,2026-01-02,oops,002,\n,A,,,003,\n'


def workbook():
    from openpyxl import Workbook
    out = io.BytesIO()
    wb = Workbook()
    wb.active.title = 'First'
    wb.active.append(['value','label'])
    wb.active.append([1,'A'])
    wb.active.append([2,None])
    other = wb.create_sheet('Second')
    other.append(['value'])
    other.append([99])
    wb.save(out)
    return out.getvalue()


def parquet():
    import pyarrow as pa
    import pyarrow.parquet as pq
    out = io.BytesIO()
    pq.write_table(pa.table({'value':[1., None, 3.], 'category':['A','A','B']}),out)
    return out.getvalue()


class DatasetTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name).resolve() / 'data'
        p = patch.object(config, 'DEPLOY_MODE', False)
        p.start()
        self.addCleanup(p.stop)
        self.registry = Registry(self.root)

    def test_profile_missing_ambiguity_statistics(self):
        table = parser.parse('table.csv', b'\xef\xbb\xbf'+CSV)
        info = profile(table)
        self.assertEqual((info['row_count'],info['column_count'],info['duplicate_rows']), (4,6,1))
        cols = {c['name']:c for c in info['columns']}
        self.assertEqual(cols['number']['type'],'numeric')
        self.assertEqual(cols['group']['type'],'categorical')
        self.assertEqual(cols['date']['type'],'temporal')
        self.assertEqual(cols['empty']['nulls'],4)
        self.assertTrue(cols['mixed']['ambiguities'])
        self.assertTrue(cols['identifier']['ambiguities'])
        stats = describe(table,['number','group','mixed'])['statistics']
        self.assertAlmostEqual(stats['number']['mean'],7/3)
        self.assertEqual(stats['number']['median'],3)
        self.assertEqual(stats['number']['quantiles']['0.25'],2)
        self.assertEqual(stats['number']['missing_fraction'],.25)
        self.assertNotIn('mean',stats['mixed'])
        self.assertEqual(stats['group']['category_counts'], {'str:A':2,'str:B':2})

    def test_versions_dedup_metadata_stats_restart(self):
        r, created = self.registry.register('Study','study.csv',CSV)
        self.assertTrue(created)
        duplicate, created = self.registry.register('Renamed','other.csv',CSV)
        self.assertFalse(created)
        self.assertEqual(r['id'],duplicate['id'])
        changed, _ = self.registry.register('Study','study.csv',CSV.replace(b'1,A',b'2,A'))
        self.assertEqual(changed['version'],2)
        self.assertNotEqual(r['sha256'],changed['sha256'])
        dictionary = r['dictionary']
        dictionary['number'] = dict(description='Synthetic measurement', units='kg')
        self.registry.dictionary(r['id'],dictionary)
        self.assertEqual(self.registry.source(r),CSV)
        saved = self.registry.statistics(r['id'],['number','group'])
        self.assertEqual(saved['dataset_sha256'],hashlib.sha256(CSV).hexdigest())
        self.assertEqual(saved['parameters']['std_ddof'],1)
        reopened = Registry(self.root)
        self.assertEqual(reopened.saved_statistics(r['id']),[saved])
        self.assertEqual(reopened.get(r['id'])['dictionary'],dictionary)
        self.assertEqual(reopened.reprofile(r['id'])['row_count'],4)
        self.assertEqual(len(reopened.inventory('study')),2)
        self.assertEqual(len(reopened.inventory("' OR 1=1")),0)
        with self.assertRaises(ValueError):
            reopened.dictionary(r['id'],{'injected':{}})
        with self.assertRaises(ValueError):
            reopened.statistics(r['id'],['__import__("os")'])

    def test_xlsx_sheet_identity_options_and_padding(self):
        content = workbook()
        manifest = parser.parse('book.xlsx',content)
        self.assertEqual(manifest['sheets'],['First','Second'])
        with self.assertRaises(ValueError):
            self.registry.register('Book','book.xlsx',content)
        r, _ = self.registry.register('Book','book.xlsx',content,'First')
        r2, _ = self.registry.register('Book','book.xlsx',content,'Second')
        self.assertEqual(r2['version'],2)
        self.assertEqual(r['sha256'],r2['sha256'])
        self.assertEqual(r['profile']['columns'][1]['nulls'],1)
        self.assertEqual(self.registry.inspect(r2['id'])[1]['rows'],[[99]])
        with self.assertRaises(ValueError):
            parser.parse('book.xlsx',content,'Absent')

    def test_parquet_scalar_types_missingness(self):
        r, _ = self.registry.register('Arrow','arrow.parquet',parquet())
        self.assertEqual(r['profile']['columns'][0]['nulls'],1)
        self.assertEqual(self.registry.statistics(r['id'],['value'])['statistics']['value']['mean'],2)
        self.assertIn('double',r['profile']['physical_types'])
        import pyarrow as pa
        import pyarrow.parquet as pq
        out = io.BytesIO()
        pq.write_table(pa.table({'nested':[[1],[2]]}),out)
        with self.assertRaisesRegex(ValueError,'scalar'):
            parser.parse('bad.parquet',out.getvalue())

    def test_invalid_files_and_bounds(self):
        for name, content in [('bad.csv',b'a\n\xff'),('bad.csv',b'a,b\n1\n'),('bad.csv',b'a,a\n1,2'),('bad.xlsx',b'not zip'),('bad.parquet',b'PAR1badPAR1'),('bad.csv',workbook()),('bad.csv',b'a\n\x00'),('../bad.csv',CSV),('C:\\bad.csv',CSV),('bad.exe',CSV),('empty.csv',b'')]:
            with self.subTest(name=name,content=content[:12]), self.assertRaises(ValueError):
                parser.parse(name,content)
        for content in [b'a\n'+b'x'*4097, (','.join('c'+str(i) for i in range(101))+'\n').encode(), b'a\n'+b'1\n'*20001]:
            with self.assertRaises(ValueError):
                parser.parse('big.csv',content)
        with self.assertRaisesRegex(ValueError,'10 MiB'):
            parser.parse('big.csv',b'x'*(parser.MAX_BYTES+1))
        with self.assertRaisesRegex(ValueError,'UTF-8'):
            parser.parse('latin.csv',b'a\n\xe9')
        # Dangerous-looking strings remain inert data.
        table = parser.parse('safe.csv',b'a\n=cmd()\n__import__(os)\n')
        self.assertEqual(table['rows'][0][0],'=cmd()')

    def test_zip_bomb_macros_and_external_links(self):
        for entry, content in [('xl/vbaProject.bin',b'code'),('xl/externalLinks/externalLink1.xml',b'link'),('../evil',b'path'),('big',b'a'*100000)]:
            out = io.BytesIO()
            with ZipFile(out,'w',compression=8) as z:
                z.writestr('[Content_Types].xml','x')
                z.writestr('xl/workbook.xml','x')
                z.writestr(entry,content)
            with self.assertRaises(ValueError):
                parser.parse('evil.xlsx',out.getvalue())

    def test_permissions_integrity_symlinks_demo(self):
        r,_ = self.registry.register('Private','private.csv',CSV)
        self.assertEqual(self.registry.root.stat().st_mode & 0o777,0o700)
        self.assertEqual(self.registry.db.stat().st_mode & 0o777,0o600)
        blob = self.registry._blob(r['sha256'])
        self.assertEqual(blob.stat().st_mode & 0o777,0o400)
        blob.chmod(0o600)
        blob.write_bytes(b'tampered')
        with self.assertRaisesRegex(ValueError,'integrity'):
            self.registry.inspect(r['id'])
        blob.unlink()
        blob.symlink_to(self.registry.db)
        with self.assertRaises(ValueError):
            self.registry.inspect(r['id'])
        outside = Path(self.temp.name).resolve()/'outside'
        outside.mkdir()
        redirected = Path(self.temp.name).resolve()/'redirect'
        redirected.symlink_to(outside, target_is_directory=True)
        with self.assertRaisesRegex(ValueError,'symbolic'):
            Registry(redirected)
        self.registry.db.unlink()
        self.registry.db.symlink_to(outside/'database')
        with self.assertRaises(ValueError):
            self.registry.inventory()
        with patch.object(config,'DEPLOY_MODE',True):
            with self.assertRaisesRegex(ValueError,'read-only'):
                self.registry.source(r)
            with self.assertRaisesRegex(ValueError,'read-only'):
                self.registry.inventory()
        with patch.object(config,'DEPLOY_MODE',True), self.assertRaisesRegex(ValueError,'read-only'):
            Registry(Path(self.temp.name).resolve()/'never-created')
        self.assertFalse((Path(self.temp.name).resolve()/'never-created').exists())
        with self.assertRaises(ValueError):
            Registry(config.ROOT/'demo')

    def test_failure_recovery_no_partial_records(self):
        with self.registry.connection() as db:
            db.execute("CREATE TRIGGER fail_insert BEFORE INSERT ON datasets BEGIN SELECT RAISE(ABORT,'synthetic failure'); END")
        with self.assertRaises(sqlite3.IntegrityError):
            self.registry.register('Retry','retry.csv',CSV)
        self.assertEqual(self.registry.inventory(),[])
        with self.registry.connection() as db:
            db.execute('DROP TRIGGER fail_insert')
        record,created = self.registry.register('Retry','retry.csv',CSV)
        self.assertTrue(created)
        self.assertEqual(record['version'],1)
        self.assertEqual(self.registry.source(record),CSV)

    def test_concurrent_import_deduplicates(self):
        with concurrent.futures.ThreadPoolExecutor(max_workers=2) as executor:
            results = list(executor.map(lambda _: self.registry.register('Race','race.csv',CSV), range(2)))
        self.assertEqual(sum(created for _,created in results),1)
        self.assertEqual(len(self.registry.inventory()),1)

    def test_inventory_pages_unicode_search_and_reproducibility(self):
        from dataset_registry import now
        record, _ = self.registry.register('Étude', 'test.csv', CSV)
        first = self.registry.statistics(record['id'], ['number'])
        second = self.registry.statistics(record['id'], ['number'])
        for key in ('id','timestamp'):
            first.pop(key)
            second.pop(key)
        self.assertEqual(first, second)
        self.assertEqual(len(self.registry.inventory('ÉTUDE')), 1)
        # Synthetic registry fixtures isolate pagination from parser throughput.
        with self.registry.connection() as db:
            for i in range(101):
                db.execute('INSERT INTO datasets SELECT ?,name,?,filename,format,?,bytes,?,origin,options,profile,dictionary FROM datasets WHERE id=?',
                           (f'fixture-{i}',i+2,f'{i:064x}',now(),record['id']))
        self.assertEqual(self.registry.inventory_count(), 102)
        self.assertEqual(len(self.registry.inventory()), 100)
        self.assertEqual(len(self.registry.inventory(page=1)), 2)
        self.assertFalse({r['id'] for r in self.registry.inventory()} & {r['id'] for r in self.registry.inventory(page=1)})

    def test_atomic_source_failure_can_retry(self):
        import dataset_registry
        with patch.object(dataset_registry.os, 'fsync', side_effect=OSError('synthetic disk failure')):
            with self.assertRaises(OSError):
                self.registry.register('Atomic', 'atomic.csv', CSV)
        self.assertEqual(self.registry.inventory(), [])
        self.assertEqual(list(self.registry.root.glob('*.source')), [])
        self.assertEqual(list(self.registry.root.glob('.source-*')), [])
        record, _ = self.registry.register('Atomic', 'atomic.csv', CSV)
        self.assertEqual(self.registry.source(record), CSV)

    def test_worker_timeout_kills_process(self):
        import subprocess
        from unittest.mock import Mock
        worker = Mock()
        worker.communicate.side_effect = [subprocess.TimeoutExpired('synthetic', 30), (None,None)]
        with patch.object(parser.subprocess, 'Popen', return_value=worker):
            with self.assertRaisesRegex(ValueError, '30 seconds'):
                parser.parse('slow.csv', CSV)
        worker.kill.assert_called_once()
        self.assertEqual(worker.communicate.call_count, 2)

    def test_remaining_resource_ceilings(self):
        for content in [
            ((','.join('c'+str(i) for i in range(100))+'\n') + (','.join('1' for _ in range(100))+'\n')*1001).encode(),
            b'a,b\n' + (b'x'*4096+b','+b'x'*4096+b'\n')*250,
        ]:
            with self.assertRaises(ValueError):
                parser.parse('bounded.csv', content)
        from openpyxl import Workbook
        wb = Workbook()
        for i in range(20):
            wb.create_sheet(str(i))
        out = io.BytesIO()
        wb.save(out)
        with self.assertRaisesRegex(ValueError, '20 sheets'):
            parser.parse('sheets.xlsx', out.getvalue())
        import pyarrow as pa
        import pyarrow.parquet as pq
        for table in [pa.table({'a': list(range(20001))}), pa.table({'a': ['x'*4097]}), pa.table({'a':['x'*4096]*9000})]:
            out = io.BytesIO()
            pq.write_table(table,out,use_dictionary=False)
            with self.assertRaises(ValueError):
                parser.parse('bounded.parquet',out.getvalue())

    def test_dataset_and_document_ingestion_share_workspace(self):
        import document_ingestion as documents
        import refresh
        import workspace
        for folder in config.PARA_FOLDERS:
            (self.root/'wiki'/folder).mkdir(parents=True)
        patches = [(config,'DATA_DIR',self.root),(config,'EMBEDDING_PROVIDER','ollama'),
                   (workspace,'DATA_DIR',self.root),(workspace,'DEPLOY_MODE',False),
                   (workspace,'ENABLE_CAPTURE',True),(workspace,'ENABLE_REFRESH_BRAIN',True)]
        from contextlib import ExitStack
        with ExitStack() as stack:
            for module,name,value in patches:
                stack.enter_context(patch.object(module,name,value))
            docs = documents.save_and_index('source.csv',b'Name,Value\nAlpha,42',embedder=lambda text: [1.,0.,0.])
            pointer = (self.root/'knowledge/current.json').read_bytes()
            with patch.object(refresh, 'refresh', side_effect=AssertionError('Datasets must not index')):
                record, _ = self.registry.register('Structured', 'structured.csv', CSV)
                self.registry.statistics(record['id'], ['number'])
            self.assertEqual((self.root/'knowledge/current.json').read_bytes(), pointer)
            self.assertEqual(documents.history()[0]['id'],docs['id'])
            self.assertTrue(any(n.frontmatter.get('document') for n in workspace.notes()))
            self.assertEqual(self.registry.source(record),CSV)

    def test_demo_ui_never_opens_registry(self):
        from streamlit.testing.v1 import AppTest
        import dataset_ui
        with patch.object(config,'DEPLOY_MODE',True), patch.object(dataset_ui,'Registry',side_effect=AssertionError('private read')):
            at = AppTest.from_string("from dataset_ui import render; render()").run()
            self.assertFalse(at.exception)
            self.assertIn('read-only',at.info[0].value)


if __name__ == '__main__':
    unittest.main()
