"""Bounded local document extraction and recoverable, transactional imports."""
from __future__ import annotations
import csv
import hashlib
import io
import json
import os
from pathlib import Path
import re
import tempfile
from datetime import datetime, timezone
from zipfile import ZipFile
from urllib.parse import urlparse
import ipaddress
import yaml
import config
import refresh as pipeline

MAX_BYTES = 25 * 1024 * 1024
MAX_TEXT = 1_000_000
MAX_PAGES = 200
MAX_SHEETS = 20
MAX_ROWS = 10_000
MAX_CELLS = 100_000
MAX_SEGMENTS = 1000
CHUNK = 1000


def citation(meta):
    if not meta:
        return ''
    location = f"page {meta['page']}" if 'page' in meta else f"sheet {meta['sheet']}, rows {meta['row_start']}–{meta['row_end']}"
    return f"{meta['filename']} · {location}"


def extract(filename, content):
    """Pure preview: never persists source bytes or calls a model."""
    if not content or len(content) > MAX_BYTES:
        raise ValueError('File is empty or exceeds 25 MiB.')
    name = re.sub(r'[\x00-\x1f\x7f]', '_', filename.replace('\\', '/').split('/')[-1])[:180]
    suffix = Path(name).suffix.lower()
    if suffix not in ('.pdf', '.csv', '.xlsx'):
        raise ValueError('Supported formats: PDF, UTF-8 CSV, XLSX.')
    sha = hashlib.sha256(content).hexdigest()
    result = dict(filename=name, sha256=sha, source_type=suffix[1:], warnings=[], segments=[])
    total = 0
    def add(text, **location):
        nonlocal total
        total += len(text)
        if total > MAX_TEXT:
            raise ValueError('Extracted text exceeds 1,000,000 characters.')
        for start in range(0, len(text), CHUNK):
            if len(result['segments']) >= MAX_SEGMENTS:
                raise ValueError('Document exceeds 1,000 segments.')
            result['segments'].append(dict(text=text[start:start+CHUNK], **location))
    def rows(values, sheet):
        cells = 0
        for number, row in enumerate(values, 1):
            cells += len(row)
            if number > MAX_ROWS or cells > MAX_CELLS:
                raise ValueError('Sheet exceeds 10,000 rows or 100,000 cells.')
            text = ' | '.join('' if value is None else str(value) for value in row)
            if text.strip(' |'):
                add(text, sheet=sheet, row_start=number, row_end=number)
    try:
        if suffix == '.pdf':
            if not content.startswith(b'%PDF-'):
                raise ValueError('Invalid PDF signature.')
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(content), strict=True)
            if reader.is_encrypted:
                raise ValueError('Encrypted PDFs are unsupported.')
            if len(reader.pages) > MAX_PAGES:
                raise ValueError('PDF exceeds 200 pages.')
            result['pages'] = len(reader.pages)
            for number, page in enumerate(reader.pages, 1):
                text = page.extract_text() or ''
                if not text.strip():
                    result['warnings'].append(f'Page {number}: blank/image-only; OCR unsupported.')
                else:
                    add(text, page=number)
        elif suffix == '.csv':
            text = content.decode('utf-8-sig', errors='strict')
            if '\x00' in text or content.startswith((b'PK', b'%PDF-')):
                raise ValueError('Invalid CSV content.')
            csv.field_size_limit(MAX_TEXT)
            rows(csv.reader(io.StringIO(text, newline=''), strict=True), 'CSV')
            result['sheets'] = ['CSV']
        else:
            if not content.startswith(b'PK\x03\x04'):
                raise ValueError('Invalid XLSX signature.')
            with ZipFile(io.BytesIO(content)) as archive:
                entries = archive.infolist()
                if len(entries) > 2000 or sum(e.file_size for e in entries) > 50 * 1024 * 1024:
                    raise ValueError('XLSX decompressed size exceeds limits.')
                for entry in entries:
                    if entry.flag_bits & 1 or entry.file_size > 20 * 1024 * 1024 or entry.file_size > max(1, entry.compress_size) * 200:
                        raise ValueError('Unsafe XLSX compression or encryption.')
                names = archive.namelist()
                if '[Content_Types].xml' not in names or 'xl/workbook.xml' not in names or any('vbaproject' in n.lower() for n in names):
                    raise ValueError('Invalid XLSX workbook or macros present.')
            from openpyxl import load_workbook
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_links=False)
            try:
                if len(book.sheetnames) > MAX_SHEETS:
                    raise ValueError('Workbook exceeds 20 sheets.')
                result['sheets'] = book.sheetnames
                result['warnings'].append('Formulas are not executed; only cached values are read. External links are not followed.')
                for sheet in book:
                    if (sheet.max_row or 0) > MAX_ROWS or (sheet.max_column or 0) > MAX_CELLS:
                        raise ValueError('Sheet dimensions exceed limits.')
                    rows(sheet.iter_rows(values_only=True), sheet.title)
            finally:
                book.close()
    except UnicodeDecodeError as exc:
        raise ValueError('CSV must use UTF-8 or UTF-8 BOM encoding.') from exc
    except (ValueError, ImportError):
        raise
    except Exception as exc:
        raise ValueError('Malformed or unsupported document.') from exc
    if not result['segments']:
        raise ValueError('No extractable text. Blank/scanned documents require OCR (unsupported).')
    for ordinal, segment in enumerate(result['segments'], 1):
        segment.update(ordinal=ordinal)
    result['id'] = hashlib.sha256((sha + '\0' + name).encode()).hexdigest()
    return result


def writable():
    from workspace import writable as check
    check()
    host = urlparse(config.OLLAMA_BASE_URL).hostname or ''
    try:
        local_host = ipaddress.ip_address(host).is_loopback
    except ValueError:
        local_host = host.lower() == 'localhost'
    if not local_host:
        raise ValueError('Document indexing requires a loopback Ollama endpoint.')
    if config.DEPLOY_MODE or config.EMBEDDING_PROVIDER != 'ollama':
        raise ValueError('Document import requires local mode and Ollama embeddings.')


def atomic_json(path, value):
    fd, name = tempfile.mkstemp(dir=path.parent, prefix='.state-')
    try:
        with os.fdopen(fd, 'w') as stream:
            json.dump(value, stream, ensure_ascii=False)
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(name, path)
        pipeline.sync_directory(path.parent)
    finally:
        Path(name).unlink(missing_ok=True)


def history():
    root = config.DATA_DIR / 'imports'
    if root.is_symlink():
        raise ValueError('Unsafe import directory.')
    return [json.loads(p.read_text()) for p in sorted(root.glob('*/state.json'))
            if not p.is_symlink() and not p.parent.is_symlink()]


def save_and_index(filename=None, content=None, *, ident=None, embedder=None):
    writable()
    root = config.DATA_DIR
    for path in (root, root / 'imports', root / 'wiki', root / 'wiki/resources', root / 'knowledge'):
        if path.is_symlink():
            raise ValueError('Symlink workspace paths are unsupported.')
    with pipeline.writer_lock(root / 'knowledge'):
        imports = root / 'imports'
        if imports.is_symlink() or (root / 'wiki/resources').is_symlink():
            raise ValueError('Unsafe import directory.')
        imports.mkdir(exist_ok=True)
        if ident is None:
            preview = extract(filename, content)  # revalidate confirmation bytes
            ident = preview['id']
            bundle = imports / ident
            if not bundle.exists():
                staging = Path(tempfile.mkdtemp(dir=imports, prefix='.stage-'))
                pipeline.durable_write(staging / 'source', content)
                atomic_json(staging / 'state.json', dict(preview, status='staged', created=datetime.now(timezone.utc).isoformat()))
                pipeline.sync_directory(staging)
                os.replace(staging, bundle)
                pipeline.sync_directory(imports)
        if not re.fullmatch('[a-f0-9]{64}', ident):
            raise ValueError('Invalid import identity.')
        bundle = imports / ident
        if any(p.is_symlink() for p in (bundle, bundle / 'source', bundle / 'state.json', bundle / 'notes')):
            raise ValueError('Unsafe bundle.')
        state = json.loads((bundle / 'state.json').read_text())
        preview = extract(state['filename'], (bundle / 'source').read_bytes())
        if preview['id'] != ident:
            raise ValueError('Source checksum mismatch.')
        destination = root / 'wiki/resources' / ('document-' + ident)
        if destination.is_symlink():
            raise ValueError('Unsafe segment directory.')
        if state['status'] == 'indexed':
            return state
        notes = bundle / 'notes'
        if not notes.exists() and not destination.exists():
            pending_notes = Path(tempfile.mkdtemp(dir=bundle, prefix=".notes-"))
            for segment in preview['segments']:
                provenance = {k: v for k, v in segment.items() if k != 'text'}
                provenance.update({k: preview[k] for k in ('filename', 'sha256', 'source_type', 'warnings')})
                nid = ident + '-' + str(segment['ordinal'])
                title = citation(provenance)
                metadata = dict(id=nid, para='resources', summary=segment['text'][:300], tags=['document'], links=[], source_type=preview['source_type'], document=provenance, created=state['created'])
                pipeline.durable_write(pending_notes / (nid + '.md'), ('---\n' + yaml.safe_dump(metadata, allow_unicode=True) + '---\n# ' + title + '\n\n' + segment['text'] + '\n').encode())
            pipeline.sync_directory(pending_notes)
            os.replace(pending_notes, notes)
            pipeline.sync_directory(bundle)
        destination.parent.mkdir(parents=True, exist_ok=True)
        if notes.exists():
            os.replace(notes, destination)
            pipeline.sync_directory(destination.parent)
            pipeline.sync_directory(bundle)
        try:
            _, vectors, excluded = pipeline.verified_active_vectors(root, provider='ollama')
            if excluded or not all(ident + '-' + str(seg['ordinal']) in vectors for seg in preview['segments']):
                pipeline._refresh_locked(root, provider='ollama', embedder=embedder)
            generation = pipeline.current(root)[0].name
        except Exception:
            try:
                _, vectors, excluded = pipeline.verified_active_vectors(root, provider='ollama')
                published = not excluded and all(ident + '-' + str(seg['ordinal']) in vectors for seg in preview['segments'])
            except Exception:
                published = False
            if published:
                state.update(status='indexed', error=None, generation=pipeline.current(root)[0].name)
                atomic_json(bundle / 'state.json', state)
                return state
            os.replace(destination, notes)
            pipeline.sync_directory(destination.parent)
            pipeline.sync_directory(bundle)
            state.update(status='error', error='Indexing failed. Check local Ollama and refresh status, then retry.')
            atomic_json(bundle / 'state.json', state)
            raise ValueError(state['error']) from None
        state.update(status='indexed', error=None, generation=generation)
        atomic_json(bundle / 'state.json', state)
        return state
