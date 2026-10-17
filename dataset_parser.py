"""Bounded local table parsing. Uploaded values are never evaluated."""
from __future__ import annotations
import csv
import datetime as dt
import io
import json
import math
import os
from pathlib import Path
import subprocess
import sys
import tempfile
from zipfile import ZipFile

MAX_BYTES = 10 * 1024 * 1024
MAX_ROWS = 20_000
MAX_COLUMNS = 100
MAX_CELLS = 100_000
MAX_CELL = 4096
MAX_TEXT = 2_000_000
MAX_EXPANDED = 32 * 1024 * 1024
MAX_SHEETS = 20
MAX_OUTPUT = 16 * 1024 * 1024
TIMEOUT = 30


class ValidationError(ValueError):
    """A validation message that is safe to show without parser diagnostics."""


def filename(value):
    if not isinstance(value, str) or not value or len(value) > 180 or any(ord(c) < 32 or ord(c) == 127 for c in value) or '/' in value or '\\' in value or value in ('.', '..'):
        raise ValueError('Use a plain filename of at most 180 characters, without paths or control characters.')
    suffix = Path(value).suffix.lower()
    if suffix not in ('.csv', '.xlsx', '.parquet'):
        raise ValueError('Supported formats: UTF-8 CSV, XLSX, Parquet.')
    return suffix[1:]


def parse(name, content, sheet=None):
    """Pure preview, isolated from app state. Bound input, runtime and output."""
    fmt = filename(name)
    if not content or len(content) > MAX_BYTES:
        raise ValueError('File is empty or exceeds 10 MiB. Split it before uploading.')
    if sheet is not None and (not isinstance(sheet, str) or len(sheet) > 100):
        raise ValueError('Invalid worksheet name.')
    with tempfile.TemporaryFile() as output:
        process = subprocess.Popen([sys.executable, '-I', str(Path(__file__).resolve()), '--worker', fmt, sheet or ''], stdin=subprocess.PIPE, stdout=output, stderr=subprocess.DEVNULL, env={**{k:v for k,v in os.environ.items() if k in ('PATH','SYSTEMROOT','WINDIR','TMPDIR','LANG','LC_ALL')}, 'OPENBLAS_NUM_THREADS':'1', 'OMP_NUM_THREADS':'1', 'ARROW_NUM_THREADS':'1'})
        try:
            process.communicate(content, timeout=TIMEOUT)
        except subprocess.TimeoutExpired:
            process.kill()
            process.communicate()
            raise ValueError('Parsing exceeded 30 seconds. Split or simplify the file.') from None
        output.seek(0)
        raw = output.read(MAX_OUTPUT + 1)
    if process.returncode or len(raw) > MAX_OUTPUT:
        raise ValueError('Parser resource limit or invalid file. Split or re-export the file.')
    try:
        result = json.loads(raw)
    except (ValueError, UnicodeError):
        raise ValueError('Parser could not read the file. Re-export a supported table.') from None
    if 'error' in result:
        raise ValueError(result['error'])
    return result


def _xlsx(content, selected):
    from openpyxl import load_workbook
    with ZipFile(io.BytesIO(content)) as archive:
        entries = archive.infolist()
        if len(entries) > 1000 or sum(e.file_size for e in entries) > MAX_EXPANDED:
            raise ValidationError('XLSX exceeds 1000 entries or 32 MiB expanded size.')
        names = archive.namelist()
        if len(set(names)) != len(names) or '[Content_Types].xml' not in names or 'xl/workbook.xml' not in names:
            raise ValidationError('Invalid XLSX workbook.')
        for e in entries:
            lower = e.filename.lower()
            if e.flag_bits & 1 or e.file_size > MAX_EXPANDED or e.file_size > max(1, e.compress_size) * 200:
                raise ValidationError('Unsafe XLSX compression/encryption.')
            if '..' in Path(e.filename).parts or e.filename.startswith('/') or 'vbaproject' in lower or 'externallinks/' in lower:
                raise ValidationError('XLSX macros, external links and unsafe paths are unsupported.')
        if b'macroenabled' in archive.read('[Content_Types].xml').lower():
            raise ValidationError('Macro-enabled XLSX workbooks are unsupported.')
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_links=False)
    try:
        sheets = workbook.sheetnames
        if not sheets or len(sheets) > MAX_SHEETS:
            raise ValidationError('XLSX must contain 1–20 sheets.')
        if not selected:
            return dict(sheets=sheets, selection_required=True)
        if selected not in sheets:
            raise ValidationError('Select an existing worksheet.')
        ws = workbook[selected]
        if (ws.max_row or 0) > MAX_ROWS + 1 or (ws.max_column or 0) > MAX_COLUMNS:
            raise ValidationError('Worksheet dimensions exceed 20,000 data rows or 100 columns.')
        # Ignore incorrect undersized producer dimensions; validate actual rows below.
        ws.reset_dimensions()
        return _table(ws.iter_rows(values_only=True), pad_rows=True, sheets=sheets, sheet=selected)
    finally:
        workbook.close()


def _parquet(content):
    import pyarrow as pa
    import pyarrow.parquet as pq
    if not content.startswith(b'PAR1') or not content.endswith(b'PAR1'):
        raise ValidationError('Invalid Parquet signature.')
    pf = pq.ParquetFile(pa.BufferReader(content), thrift_string_size_limit=1_000_000, thrift_container_size_limit=100_000)
    meta = pf.metadata
    if meta.num_rows > MAX_ROWS or meta.num_columns > MAX_COLUMNS or meta.num_rows * meta.num_columns > MAX_CELLS or meta.num_row_groups > 100:
        raise ValidationError('Parquet exceeds row, column, cell or 100 row-group limit.')
    if sum(meta.row_group(i).total_byte_size for i in range(meta.num_row_groups)) > MAX_EXPANDED:
        raise ValidationError('Parquet exceeds 32 MiB decoded size.')
    for field in pf.schema_arrow:
        t = field.type
        if not (pa.types.is_integer(t) or pa.types.is_floating(t) or pa.types.is_decimal(t) or pa.types.is_boolean(t) or pa.types.is_string(t) or pa.types.is_large_string(t) or pa.types.is_date(t) or pa.types.is_timestamp(t)):
            raise ValidationError('Parquet supports only scalar numeric, boolean, string, date and timestamp columns.')
    def rows():
        yield pf.schema_arrow.names
        for batch in pf.iter_batches(batch_size=256, use_threads=False):
            cols = [c.to_pylist() for c in batch.columns]
            yield from zip(*cols)
    return _table(rows(), physical_types=[str(f.type) for f in pf.schema_arrow])


def _table(values, pad_rows=False, **metadata):
    iterator = iter(values)
    try:
        header = list(next(iterator))
    except StopIteration:
        raise ValidationError('Table needs a header row.') from None
    if not 1 <= len(header) <= MAX_COLUMNS:
        raise ValidationError('Table must contain 1–100 columns.')
    names = [str(v).strip() if v is not None else '' for v in header]
    if any(not n or len(n) > 128 for n in names) or len(set(names)) != len(names):
        raise ValidationError('Column names must be unique, nonempty and at most 128 characters.')
    rows, total = [], 0
    for row in iterator:
        row = list(row)
        if pad_rows and len(row) < len(names):
            row.extend([None] * (len(names)-len(row)))
        if len(row) != len(names):
            raise ValidationError('Every row must match the header width; check delimiters or worksheet layout.')
        if len(rows) >= MAX_ROWS or (len(rows) + 1) * len(names) > MAX_CELLS:
            raise ValidationError('Table exceeds 20,000 rows or 100,000 cells.')
        clean = []
        for v in row:
            if isinstance(v, (dt.datetime, dt.date)):
                v = v.isoformat()
            elif v is not None and not isinstance(v, (str, int, float, bool)):
                v = str(v)
            if isinstance(v, float) and not math.isfinite(v):
                v = None
            if isinstance(v, str):
                if len(v) > MAX_CELL or any(ord(c) < 32 and c not in '\t\r\n' for c in v):
                    raise ValidationError('Cell exceeds 4096 characters or contains unsupported control characters.')
                if not v.strip():
                    v = None
            total += len(str(v)) if v is not None else 0
            if total > MAX_TEXT:
                raise ValidationError('Table exceeds 2,000,000 decoded value characters.')
            clean.append(v)
        rows.append(clean)
    return dict(columns=names, rows=rows, **metadata)


def _worker(fmt, sheet):
    # OS limits supplement application bounds; this is not an adversarial sandbox.
    try:
        import resource
        resource.setrlimit(resource.RLIMIT_CPU, (20, 20))
        resource.setrlimit(resource.RLIMIT_FSIZE, (MAX_OUTPUT, MAX_OUTPUT))
        if sys.platform.startswith('linux'):
            resource.setrlimit(resource.RLIMIT_AS, (2 * 1024**3, 2 * 1024**3))
    except (ImportError, ValueError, OSError):
        pass
    try:
        content = sys.stdin.buffer.read(MAX_BYTES + 1)
        if not content or len(content) > MAX_BYTES:
            raise ValidationError('File is empty or exceeds 10 MiB.')
        if fmt == 'csv':
            try:
                text = content.decode('utf-8-sig', errors='strict')
            except UnicodeError:
                raise ValidationError('Unsupported CSV encoding. Re-export as UTF-8 (BOM accepted).') from None
            if '\x00' in text or content.startswith((b'PK', b'PAR1', b'%PDF')):
                raise ValidationError('Content does not match CSV extension.')
            csv.field_size_limit(MAX_CELL)
            result = _table(csv.reader(io.StringIO(text, newline=''), strict=True))
        elif fmt == 'xlsx':
            result = _xlsx(content, sheet)
        else:
            result = _parquet(content)
        print(json.dumps(result, allow_nan=False))
    except Exception as exc:
        # Parser details may contain private cell values: expose only our validation messages.
        message = str(exc) if isinstance(exc, ValidationError) else 'Corrupt or unsupported file. Re-export as CSV, XLSX or scalar Parquet.'
        print(json.dumps({'error': message}))


if __name__ == '__main__':
    _worker(sys.argv[2], sys.argv[3])
